#!/usr/bin/env python3
"""Extract every page/slide of the course files into curriculum/extracted/ (not committed).

For each source file this writes:
  <slug>/page-001.txt …   the slide/page text (plus speaker notes for .pptx)
  <slug>/page-001.png …   a rendered image of each PDF page (for diagrams and formulas)
  index.json              page counts and the first line of every page, for the source map

Needs PyMuPDF (`pip install pymupdf`) and python-pptx (`pip install python-pptx`).
"""
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'source_materials'
OUT = ROOT / 'curriculum' / 'extracted'

EXPECTED = {
    'Review_Quiz1.pdf': 'review',
    'Lecture01&02 (3).pdf': 'L01-02',
    'Lecture03&04 (2).pdf': 'L03-04',
    'Lecture05&06_forupload.pptx': 'L05-06',
    'Lecture07&08 (2).pdf': 'L07-08',
    'Lecture09&10.pdf': 'L09-10',
}


def find(name):
    """Match the expected file name loosely (download suffixes like ' (3)' vary)."""
    exact = SRC / name
    if exact.exists():
        return exact
    stem = re.sub(r'\s*\(\d+\)', '', Path(name).stem).lower().replace('_forupload', '')
    for f in SRC.iterdir():
        s = re.sub(r'\s*\(\d+\)', '', f.stem).lower().replace('_forupload', '')
        if s == stem and f.suffix.lower() == Path(name).suffix.lower():
            return f
    return None


def pdf(path, dest):
    import fitz  # PyMuPDF
    doc = fitz.open(path)
    pages = []
    for i, page in enumerate(doc, 1):
        text = page.get_text('text')
        (dest / f'page-{i:03d}.txt').write_text(text, encoding='utf-8')
        page.get_pixmap(dpi=80).save(dest / f'page-{i:03d}.png')
        pages.append(text)
    return pages


def pptx(path, dest):
    from pptx import Presentation
    prs = Presentation(path)
    pages = []
    for i, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append('\n'.join(p.text for p in shape.text_frame.paragraphs if p.text.strip()))
            if getattr(shape, 'has_table', False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(' | '.join(c.text for c in row.cells))
        notes = ''
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
        text = '\n'.join(p for p in parts if p.strip())
        if notes.strip():
            text += '\n\n[SPEAKER NOTES]\n' + notes
        (dest / f'page-{i:03d}.txt').write_text(text, encoding='utf-8')
        pages.append(text)
    return pages


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    index, missing = {}, []
    for name, slug in EXPECTED.items():
        path = find(name)
        if not path:
            missing.append(name)
            continue
        dest = OUT / slug
        dest.mkdir(exist_ok=True)
        pages = pdf(path, dest) if path.suffix.lower() == '.pdf' else pptx(path, dest)
        index[slug] = {'file': path.name, 'pages': len(pages),
                       'firstLines': [next((l.strip() for l in p.splitlines() if l.strip()), '') for p in pages]}
        print(f'{slug}: {path.name} -> {len(pages)} pages')
    (OUT / 'index.json').write_text(json.dumps(index, indent=1), encoding='utf-8')
    if missing:
        print('MISSING:', ', '.join(missing))
        sys.exit(1)


if __name__ == '__main__':
    main()
